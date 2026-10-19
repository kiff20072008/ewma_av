"""
Рисунки к статье в формате «Вестника ТОГУ».

    python make_figures.py --ugr ugr16 --boun results_boun --out figures

Рис. 1. (а) периодограмма интенсивности потоков на непрерывном участке
        обучающей части UGR'16 длиной 11 недель; (б) суточный профиль.
Рис. 2. UDP-трафик магистрали кампусной сети (BOUN) и тревоги трёх методов.

Требования журнала: ширина не более 150 мм (полоса набора при полях журнала —
130 мм), 300 dpi, копия каждого рисунка в JPG, графики — с возможностью
редактирования. Поэтому, помимо PNG/JPG, создаётся figures_data.xlsx:
исходные данные каждого рисунка и редактируемые диаграммы Excel.

Оформление рассчитано на чёрно-белую печать: на рис. 2 методы различаются
подписями строк, а не только цветом; шрифт Liberation Serif метрически
совпадает с Times New Roman.
"""
from __future__ import annotations

import argparse
import warnings
from pathlib import Path

import numpy as np
import pandas as pd
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.ticker import FuncFormatter

from ewma_av.ugr16_minute import load_ugr16_minute
from ewma_av.datasets import profile_by_hour
from ewma_av.boun import schedule_labels
from ewma_av.detectors import EwmaAV, AEWMA, HampelMAD
from ewma_av import metrics as M

warnings.filterwarnings("ignore")

BLUE, ORANGE = "#2a78d6", "#eb6834"      # проверены валидатором палитры
GRAY, GRAY_LIGHT, INK, INK2 = "#8f8d88", "#e4e2dd", "#0b0b0b", "#52514e"
CM = 1 / 2.54
WIDTH = 13 * CM                           # ширина полосы набора журнала

plt.rcParams.update({
    "font.family": "Liberation Serif",
    "font.size": 8, "axes.titlesize": 8, "axes.labelsize": 8,
    "xtick.labelsize": 7, "ytick.labelsize": 7, "legend.fontsize": 7,
    "axes.edgecolor": INK2, "axes.labelcolor": INK, "xtick.color": INK2,
    "ytick.color": INK2, "axes.spines.top": False, "axes.spines.right": False,
    "axes.grid": True, "grid.color": GRAY_LIGHT, "grid.linewidth": 0.5,
    "axes.axisbelow": True, "lines.linewidth": 1.1, "savefig.dpi": 300,
    "savefig.bbox": "tight", "savefig.pad_inches": 0.02,
    "legend.frameon": False,
})
COMMA = FuncFormatter(lambda v, _: f"{v:g}".replace(".", ",").replace("-", "−"))


def panel_label(ax, s):
    ax.text(-0.16, 1.03, s, transform=ax.transAxes, fontsize=9,
            fontweight="bold", va="bottom", ha="left", color=INK)


def contiguous_segments(ts):
    d = np.diff(ts.values).astype("timedelta64[m]").astype(int)
    cuts = np.where(d > 1)[0]
    b = [0] + list(cuts + 1) + [len(ts)]
    return list(zip(b[:-1], b[1:]))


def spectrum_11_weeks(x, ts):
    a, b = max(contiguous_segments(ts), key=lambda ab: ab[1] - ab[0])
    n = 11 * 7 * 1440
    v = np.array(x[b - n:b], dtype=float)    # копия: исходный ряд не трогаем
    v = v - v.mean()
    f = np.fft.rfftfreq(n, d=60.0)[1:]
    p = (np.abs(np.fft.rfft(v)) ** 2)[1:]
    return 1 / f / 3600, 100 * p / p.sum()


def save(fig, out: Path, stem: str):
    fig.savefig(out / f"{stem}.png")
    fig.savefig(out / f"{stem}.jpg", pil_kwargs={"quality": 95})
    plt.close(fig)


def fig1(x, ts, out):
    period_h, rel = spectrum_11_weeks(x, ts)
    ph = profile_by_hour(x, ts)
    fig, (a, b) = plt.subplots(1, 2, figsize=(WIDTH, 4.3 * CM))

    keep = (period_h >= 3) & (period_h <= 1000)
    a.plot(period_h[keep], rel[keep], color=BLUE, linewidth=0.8)
    a.set_xscale("log")
    a.set_xticks([3, 6, 12, 24, 168, 720])
    a.set_xticklabels(["3", "6", "12", "24", "168", "720"])
    a.set_xlim(3, 1000)
    a.set_xlabel("Период, ч")
    a.set_ylabel("Доля мощности, %")
    for target, text, dx, dy in [(24.0, "24 ч", 4, 0), (168.0, "168 ч", 4, 0), (12.0, "12 ч", 2, 3)]:
        j = int(np.argmin(np.abs(period_h - target)))
        a.annotate(text, (period_h[j], rel[j]), xytext=(dx, dy),
                   textcoords="offset points", fontsize=7, color=INK,
                   va="bottom" if dy else "center")
    panel_label(a, "а")

    h = ph.index.to_numpy()
    b.fill_between(h, ph["q10"] / 1e3, ph["q90"] / 1e3, color=BLUE, alpha=0.16,
                   linewidth=0, label="10–90 %")
    b.plot(h, ph["median"] / 1e3, color=BLUE, marker="o", markersize=2.2,
           label="медиана")
    b.set_xticks(range(0, 24, 6))
    b.set_xlim(-0.3, 23.3)
    b.set_xlabel("Час суток")
    b.set_ylabel("Потоков в минуту, тыс.")
    b.legend(loc="upper left", handlelength=1.4)
    panel_label(b, "б")
    fig.tight_layout(w_pad=1.6)
    save(fig, out, "fig1_ugr16")
    return (pd.DataFrame({"period_h": period_h[keep], "power_pct": rel[keep]}),
            ph[["median", "q10", "q90"]].reset_index())


def fig2(counters_csv, out):
    c = pd.read_csv(counters_csv)
    x = c["udp"].to_numpy(float)
    truth, ignore = schedule_labels(len(x))
    WARMUP = 60
    keep = np.ones(len(x), bool)
    keep[:WARMUP] = False
    keep &= ~ignore
    rows = [("предлагаемая схема", EwmaAV(warmup=WARMUP), BLUE),
            ("AEWMA", AEWMA(0.2, 3.0, warmup=WARMUP), ORANGE),
            ("медиана и MAD", HampelMAD(WARMUP), GRAY)]
    alarms = {}
    for name, d, col in rows:
        s = d.score(x)
        thr = M.calibrate_threshold(s[keep], 0.01, clean_mask=~truth[keep])
        alarms[name] = ((s > thr) & keep, col)

    fig, (top, bot) = plt.subplots(
        2, 1, figsize=(WIDTH, 4.5 * CM), sharex=True,
        gridspec_kw={"height_ratios": [2.2, 1], "hspace": 0.12})
    tsec = np.arange(len(x))
    for k, (a0, a1) in enumerate([(80, 100), (180, 200), (280, 300), (380, 400)]):
        for ax in (top, bot):
            ax.axvspan(a0, a1, color=GRAY_LIGHT, linewidth=0, zorder=0)
        top.text((a0 + a1) / 2, 1.0, f"атака {k + 1}",
                 transform=top.get_xaxis_transform(), ha="center",
                 va="bottom", fontsize=7, color=INK2)
    top.plot(tsec, x / 1e3, color=INK, linewidth=0.6)
    top.set_ylabel("UDP-пакетов/с, тыс.")
    top.yaxis.set_major_formatter(COMMA)
    top.grid(axis="x", visible=False)
    names = list(alarms)
    for i, name in enumerate(names):
        p, col = alarms[name]
        y = len(names) - 1 - i
        bot.vlines(np.flatnonzero(p), y - 0.32, y + 0.32, color=col, linewidth=1.2)
    bot.set_yticks(range(len(names)), names[::-1])
    bot.set_ylim(-0.6, len(names) - 0.4)
    bot.set_xlabel("Время от начала записи, с")
    bot.grid(False)
    bot.set_xlim(0, len(x))
    bot.tick_params(axis="y", length=0)
    save(fig, out, "fig2_boun_udp")
    data = pd.DataFrame({"t_s": tsec, "udp_pkts": x, "attack": truth.astype(int),
                         "guard_band": ignore.astype(int)})
    for name in names:
        data["alarm_" + name] = alarms[name][0].astype(int)
    return data


def excel(spec, prof, boun, path):
    """Данные рисунков и редактируемые диаграммы Excel."""
    from openpyxl import Workbook
    from openpyxl.chart import ScatterChart, LineChart, Reference, Series
    from openpyxl.styles import Font

    wb = Workbook()
    f = Font(name="Times New Roman", size=11)
    fb = Font(name="Times New Roman", size=11, bold=True)

    def sheet(title, df, headers):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for c in ws[1]:
            c.font = fb
        for r in df.itertuples(index=False):
            ws.append(list(r))
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = f
        return ws

    wb.remove(wb.active)
    ws = sheet("Рис1а_спектр", spec, ["Период, ч", "Доля мощности, %"])
    ch = ScatterChart()
    ch.title = "Рис. 1, а. Периодограмма"
    ch.x_axis.title, ch.y_axis.title = "Период, ч", "Доля мощности, %"
    ch.x_axis.scaling.logBase = 10
    s = Series(Reference(ws, min_col=2, min_row=1, max_row=ws.max_row),
               Reference(ws, min_col=1, min_row=2, max_row=ws.max_row), title_from_data=True)
    s.marker.symbol = "none"
    ch.series.append(s)
    ch.legend = None
    ws.add_chart(ch, "D2")

    ws = sheet("Рис1б_профиль", prof, ["Час", "Медиана", "Квантиль 10 %", "Квантиль 90 %"])
    ch = LineChart()
    ch.title = "Рис. 1, б. Суточный профиль"
    ch.x_axis.title, ch.y_axis.title = "Час суток", "Потоков в минуту"
    ch.add_data(Reference(ws, min_col=2, max_col=4, min_row=1, max_row=ws.max_row),
                titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
    ws.add_chart(ch, "F2")

    heads = ["Время, с", "UDP-пакетов/с", "Окно атаки (1/0)", "Защитная полоса (1/0)"] + \
            [c.replace("alarm_", "Тревога: ") for c in boun.columns[4:]]
    ws = sheet("Рис2_BOUN", boun, heads)
    ch = LineChart()
    ch.title = "Рис. 2. UDP-трафик магистрали кампусной сети"
    ch.x_axis.title, ch.y_axis.title = "Время, с", "UDP-пакетов/с"
    ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=ws.max_row), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
    ch.width, ch.height = 24, 9
    ws.add_chart(ch, "J2")
    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--ugr", default="ugr16")
    ap.add_argument("--boun", default="results_boun")
    ap.add_argument("--out", default="figures")
    a = ap.parse_args()
    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    x, _, ts = load_ugr16_minute(a.ugr, split="train")
    spec, prof = fig1(x, ts, out)
    print("рис. 1 готов")
    boun = fig2(Path(a.boun) / "counters_BOUN_UDP_Anon.csv", out)
    print("рис. 2 готов")
    excel(spec, prof, boun, out / "figures_data.xlsx")
    print("figures_data.xlsx готов")


if __name__ == "__main__":
    main()
